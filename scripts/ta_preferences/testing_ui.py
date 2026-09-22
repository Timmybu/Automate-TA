"""Serve the local dashboard and expose a repeatable TA-scoring test endpoint.

This module is intentionally separate from the PeopleSoft scraper. It reads the
survey workbook, applies the course-topic map, writes inspectable output files,
and returns the rankings to the local browser UI.
"""

from __future__ import annotations

import argparse
import json
import threading
from datetime import datetime, timezone
from functools import partial
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from extract_ta_preferences import extract_workbook, write_json as write_preferences
from score_ta_preferences import (
    build_rankings,
    load_course_topic_weights,
    write_json as write_rankings,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = PROJECT_ROOT / "data" / "Test_Info.xlsx"
DEFAULT_COURSE_MAP = Path(__file__).with_name("course_topic_weights.json")
OUTPUT_DIR = PROJECT_ROOT / "outputs"
PREFERENCES_OUTPUT = OUTPUT_DIR / "ta-preferences.json"
RANKINGS_OUTPUT = OUTPUT_DIR / "ta-course-rankings.json"
PEOPLESOFT_REPORT_DIR = PROJECT_ROOT / "scripts" / "PeopleSoft Scraper" / "reports"
PEOPLESOFT_REQUIRED_HEADERS = {
    "Subject",
    "Course Number",
    "Name",
    "Class Number",
    "Type",
    "Term Enrollment",
}


def _display_path(path: Path) -> str:
    try:
        return str(path.relative_to(PROJECT_ROOT))
    except ValueError:
        return str(path)


def discover_peoplesoft_reports() -> list[Path]:
    """Return usable PeopleSoft workbooks, newest first, ignoring Excel locks."""

    if not PEOPLESOFT_REPORT_DIR.is_dir():
        return []
    reports = [
        path
        for path in PEOPLESOFT_REPORT_DIR.glob("*.xlsx")
        if not path.name.startswith("~$") and path.is_file() and path.stat().st_size > 0
    ]
    return sorted(reports, key=lambda path: path.stat().st_mtime, reverse=True)


def _course_aliases(mapped_course: str) -> set[str]:
    """Expand combined map keys such as CS 0011/0012 into report course codes."""

    subject, _, catalog_numbers = mapped_course.partition(" ")
    if not subject or not catalog_numbers:
        return {mapped_course}
    return {f"{subject} {number}" for number in catalog_numbers.split("/")}


def load_peoplesoft_report(
    report_path: Path,
    mapped_courses: list[str],
) -> dict[str, Any]:
    """Read a generated PeopleSoft workbook without changing it."""

    workbook = load_workbook(report_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ValueError(f"PeopleSoft report is empty: {report_path.name}") from exc
        headers = [str(value).strip() if value is not None else "" for value in raw_headers]
        missing = sorted(PEOPLESOFT_REQUIRED_HEADERS - set(headers))
        if missing:
            raise ValueError(
                f"PeopleSoft report {report_path.name} is missing columns: "
                + ", ".join(missing)
            )
        indexes = {header: headers.index(header) for header in headers if header}
        aliases = {
            alias: mapped_course
            for mapped_course in mapped_courses
            for alias in _course_aliases(mapped_course)
        }
        sections = []
        for raw_row in rows:
            row = list(raw_row) + [None] * (len(headers) - len(raw_row))
            subject = str(row[indexes["Subject"]] or "").strip().upper()
            raw_number = str(row[indexes["Course Number"]] or "").strip()
            if not subject or not raw_number:
                continue
            number = raw_number.zfill(4) if raw_number.isdigit() else raw_number
            course_code = f"{subject} {number}"
            enrollment = row[indexes["Term Enrollment"]]
            if not isinstance(enrollment, (int, float)):
                enrollment = None
            recitation_need = row[indexes["# Recitation TAs (PhD)"]] if "# Recitation TAs (PhD)" in indexes else None
            grader_need = row[indexes["# Graders (PhD)"]] if "# Graders (PhD)" in indexes else None
            sections.append(
                {
                    "course_code": course_code,
                    "mapped_course": aliases.get(course_code),
                    "name": str(row[indexes["Name"]] or "").strip(),
                    "class_number": str(row[indexes["Class Number"]] or "").strip(),
                    "component": str(row[indexes["Type"]] or "").strip(),
                    "enrollment": enrollment,
                    "recitation_ta_need": recitation_need,
                    "grader_need": grader_need,
                }
            )
    finally:
        workbook.close()

    return {
        "file": _display_path(report_path),
        "sheet": sheet.title,
        "section_count": len(sections),
        "mapped_section_count": sum(bool(section["mapped_course"]) for section in sections),
        "sections": sections,
    }


def run_scoring_pipeline(
    *,
    workbook_path: Path = DEFAULT_WORKBOOK,
    course_map_path: Path = DEFAULT_COURSE_MAP,
    peoplesoft_report_path: Path | None = None,
    topic_weight: float = 0.6,
    preference_weight: float = 0.4,
) -> dict[str, Any]:
    """Run extraction and scoring once, returning a browser-friendly payload."""

    if topic_weight < 0 or preference_weight < 0:
        raise ValueError("Scoring weights cannot be negative")
    if abs(topic_weight + preference_weight - 1.0) > 1e-9:
        raise ValueError("Topic and preference weights must add up to 100%")

    extracted_data = extract_workbook(workbook_path)
    course_topic_weights = load_course_topic_weights(course_map_path)
    rankings = build_rankings(
        extracted_data,
        course_topic_weights,
        topic_weight=topic_weight,
        preference_weight=preference_weight,
    )

    peoplesoft = None
    if peoplesoft_report_path is not None:
        peoplesoft = load_peoplesoft_report(
            peoplesoft_report_path,
            list(course_topic_weights),
        )
        top_by_course = {
            course: next(
                (
                    row
                    for row in rankings
                    if row["course"] == course and row["eligible"]
                ),
                None,
            )
            for course in course_topic_weights
        }
        for section in peoplesoft["sections"]:
            top = top_by_course.get(section["mapped_course"])
            section["top_candidate"] = top["name"] if top else None
            section["top_score"] = top["overall_score"] if top else None

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with PREFERENCES_OUTPUT.open("w", encoding="utf-8") as destination:
        write_preferences(extracted_data, destination)
    with RANKINGS_OUTPUT.open("w", encoding="utf-8") as destination:
        write_rankings(
            destination,
            extracted_data,
            rankings,
            course_topic_weights,
            topic_weight,
            preference_weight,
        )

    course_summaries = []
    for course in course_topic_weights:
        course_results = [row for row in rankings if row["course"] == course]
        eligible = [row for row in course_results if row["eligible"]]
        course_summaries.append(
            {
                "course": course,
                "candidate_count": len(course_results),
                "eligible_count": len(eligible),
                "top_candidate": eligible[0]["name"] if eligible else None,
                "top_score": eligible[0]["overall_score"] if eligible else None,
            }
        )

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "profile_count": extracted_data["profile_count"],
        "course_count": len(course_topic_weights),
        "ranking_count": len(rankings),
        "topic_weight": topic_weight,
        "preference_weight": preference_weight,
        "workbook": _display_path(workbook_path),
        "course_map": _display_path(course_map_path),
        "outputs": {
            "preferences": str(PREFERENCES_OUTPUT.relative_to(PROJECT_ROOT)),
            "rankings": str(RANKINGS_OUTPUT.relative_to(PROJECT_ROOT)),
        },
        "courses": course_summaries,
        "rankings": rankings,
        "peoplesoft": peoplesoft,
    }


class ScoringUIHandler(SimpleHTTPRequestHandler):
    """Serve project assets plus the local scoring API."""

    run_count = 0
    run_lock = threading.Lock()

    def _send_json(self, payload: dict[str, Any], status: HTTPStatus) -> None:
        encoded = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(encoded)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(encoded)

    def do_GET(self) -> None:  # noqa: N802 - required HTTP handler name
        if self.path == "/api/scoring-config":
            reports = discover_peoplesoft_reports()
            self._send_json(
                {
                    "workbook": str(DEFAULT_WORKBOOK.relative_to(PROJECT_ROOT)),
                    "course_map": str(DEFAULT_COURSE_MAP.relative_to(PROJECT_ROOT)),
                    "topic_weight": 0.6,
                    "preference_weight": 0.4,
                    "peoplesoft_reports": [
                        str(path.relative_to(PROJECT_ROOT)) for path in reports
                    ],
                    "peoplesoft_warning": None
                    if reports
                    else "No usable PeopleSoft .xlsx report was found. Excel lock files are ignored.",
                },
                HTTPStatus.OK,
            )
            return
        super().do_GET()

    def do_POST(self) -> None:  # noqa: N802 - required HTTP handler name
        if self.path != "/api/run-scoring":
            self._send_json({"error": "Unknown API endpoint"}, HTTPStatus.NOT_FOUND)
            return

        try:
            content_length = int(self.headers.get("Content-Length", "0"))
            if content_length > 16_384:
                raise ValueError("Request is too large")
            body = self.rfile.read(content_length)
            request = json.loads(body or b"{}")
            topic_weight = float(request.get("topic_weight", 0.6))
            preference_weight = float(request.get("preference_weight", 0.4))
            requested_report = request.get("peoplesoft_report")
            peoplesoft_report = None
            if requested_report:
                peoplesoft_report = (PROJECT_ROOT / str(requested_report)).resolve()
                report_root = PEOPLESOFT_REPORT_DIR.resolve()
                if not peoplesoft_report.is_relative_to(report_root):
                    raise ValueError("PeopleSoft report must be inside the reports folder")
                if peoplesoft_report.name.startswith("~$") or not peoplesoft_report.is_file():
                    raise ValueError("Selected PeopleSoft report is not available")
            else:
                available_reports = discover_peoplesoft_reports()
                if available_reports:
                    peoplesoft_report = available_reports[0]
            result = run_scoring_pipeline(
                peoplesoft_report_path=peoplesoft_report,
                topic_weight=topic_weight,
                preference_weight=preference_weight,
            )
            with self.run_lock:
                type(self).run_count += 1
                result["run_number"] = type(self).run_count
            self._send_json(result, HTTPStatus.OK)
        except (ValueError, OSError, json.JSONDecodeError) as exc:
            self._send_json({"error": str(exc)}, HTTPStatus.BAD_REQUEST)
        except Exception as exc:  # pragma: no cover - final boundary for the UI
            self.log_error("Scoring run failed: %s", exc)
            self._send_json(
                {"error": "The scoring run failed. Check the server terminal."},
                HTTPStatus.INTERNAL_SERVER_ERROR,
            )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run the local TA scoring test UI")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8765)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    handler = partial(ScoringUIHandler, directory=str(PROJECT_ROOT))
    server = ThreadingHTTPServer((args.host, args.port), handler)
    print(f"TA scoring test UI: http://{args.host}:{args.port}/#scoring-test")
    print("Press Ctrl+C to stop the local server.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
