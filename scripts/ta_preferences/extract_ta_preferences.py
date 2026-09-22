"""Extract TA preference survey responses from an XLSX workbook into JSON.

This script only reads and normalizes source data. It does not assign numeric values,
map topics to courses, or calculate candidate scores.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any, TextIO

from openpyxl import load_workbook


DEFAULT_SHEET = "Form Responses 1"
TOP_CHOICE_PREFIX = "Top choices ["
TOPIC_PREFIX = "Topic Areas ["

IDENTITY_HEADERS = {
    "timestamp": "Timestamp",
    "first_name": "First name",
    "last_name": "Last name",
    "email": "Pitt email address",
    "scheduling_constraints": "Scheduling constraints",
    "course_constraints": "Courses constraints",
    "goals": "Goals",
}

NOTES_HEADER_PREFIX = "Anything else we should know?"


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _serializable_timestamp(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return _text(value)


def _extract_bracketed_header(header: str, prefix: str) -> str | None:
    if header.startswith(prefix) and header.endswith("]"):
        return header[len(prefix) : -1].strip()
    return None


def _header_index(headers: list[str], expected: str) -> int:
    try:
        return headers.index(expected)
    except ValueError as exc:
        raise ValueError(f"Required column is missing: {expected!r}") from exc


def extract_workbook(
    workbook_path: str | Path,
    sheet_name: str = DEFAULT_SHEET,
) -> dict[str, Any]:
    """Return a normalized, JSON-compatible representation of the survey data."""

    path = Path(workbook_path)
    if not path.is_file():
        raise FileNotFoundError(f"Workbook not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet {sheet_name!r} was not found. "
                f"Available sheets: {', '.join(workbook.sheetnames)}"
            )

        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ValueError(f"Worksheet {sheet_name!r} is empty") from exc

        headers = [_text(value) for value in raw_headers]
        duplicates = sorted({header for header in headers if headers.count(header) > 1})
        if duplicates:
            raise ValueError(f"Duplicate worksheet headers: {', '.join(duplicates)}")

        identity_indexes = {
            field: _header_index(headers, header)
            for field, header in IDENTITY_HEADERS.items()
        }
        timestamp_index = identity_indexes["timestamp"]

        notes_indexes = [
            index
            for index, header in enumerate(headers)
            if header.startswith(NOTES_HEADER_PREFIX)
        ]
        if len(notes_indexes) != 1:
            raise ValueError(
                f"Expected one notes column beginning with {NOTES_HEADER_PREFIX!r}; "
                f"found {len(notes_indexes)}"
            )
        notes_index = notes_indexes[0]

        course_columns: dict[int, str] = {}
        topic_columns: dict[int, str] = {}
        for index, header in enumerate(headers):
            course = _extract_bracketed_header(header, TOP_CHOICE_PREFIX)
            if course is not None:
                course_columns[index] = course
                continue
            topic = _extract_bracketed_header(header, TOPIC_PREFIX)
            if topic is not None:
                topic_columns[index] = topic

        if not course_columns:
            raise ValueError(f"No {TOP_CHOICE_PREFIX!r} columns were found")
        if not topic_columns:
            raise ValueError(f"No {TOPIC_PREFIX!r} columns were found")

        profiles: list[dict[str, Any]] = []
        for raw_row in rows:
            row = list(raw_row) + [None] * (len(headers) - len(raw_row))
            if not _text(row[timestamp_index]):
                continue

            first_name = _text(row[identity_indexes["first_name"]])
            last_name = _text(row[identity_indexes["last_name"]])
            email = _text(row[identity_indexes["email"]])
            name = " ".join(
                part for part in (first_name, last_name) if part
            ).strip()

            profiles.append(
                {
                    "candidate_id": email or name,
                    "timestamp": _serializable_timestamp(row[timestamp_index]),
                    "first_name": first_name,
                    "last_name": last_name,
                    "name": name,
                    "email": email,
                    "course_preferences": {
                        course: _text(row[index])
                        for index, course in course_columns.items()
                    },
                    "topic_ratings": {
                        topic: _text(row[index])
                        for index, topic in topic_columns.items()
                    },
                    "scheduling_constraints": _text(
                        row[identity_indexes["scheduling_constraints"]]
                    ),
                    "course_constraints": _text(
                        row[identity_indexes["course_constraints"]]
                    ),
                    "goals": _text(row[identity_indexes["goals"]]),
                    "notes": _text(row[notes_index]),
                }
            )

        if not profiles:
            raise ValueError(f"Worksheet {sheet_name!r} contains no response rows")

        return {
            "schema_version": 1,
            "source": {
                "workbook": str(path),
                "sheet": sheet_name,
            },
            "courses": list(course_columns.values()),
            "topics": list(topic_columns.values()),
            "profile_count": len(profiles),
            "profiles": profiles,
        }
    finally:
        workbook.close()


def write_json(payload: dict[str, Any], destination: TextIO) -> None:
    json.dump(payload, destination, indent=2, ensure_ascii=False)
    destination.write("\n")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Extract and normalize TA preference survey responses."
    )
    parser.add_argument("workbook", type=Path, help="TA preference response workbook")
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--output", type=Path, help="Output JSON path; defaults to stdout")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    payload = extract_workbook(args.workbook, args.sheet)

    destination: TextIO
    close_destination = False
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        destination = args.output.open("w", encoding="utf-8")
        close_destination = True
    else:
        destination = sys.stdout

    try:
        write_json(payload, destination)
    finally:
        if close_destination:
            destination.close()

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
