import csv
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PLANNING_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORTS_DIR = os.path.join(PLANNING_DIR, "reports")

RECITATION_COLUMN = "# Recitation TAs (PhD)"
GRADER_COLUMN = "# Graders (PhD)"
TYPE_COLUMN = "Type"
ENROLLMENT_COLUMN = "Term Enrollment"
CAPACITY_COLUMN = "Enrollment Capacity"
PRIMARY_COMPONENTS = ("LEC", "PRA", "INT")

# csv.reader hands back plain strings, and openpyxl won't infer numeric types on its own
# -- every cell would otherwise be text, silently breaking ISNUMBER()/SUM() in Excel.
NUMERIC_COLUMNS = [CAPACITY_COLUMN, ENROLLMENT_COLUMN, RECITATION_COLUMN, GRADER_COLUMN]

RED_ENROLLMENT_CUTOFF = 20
ORANGE_ENROLLMENT_CUTOFF = 25

# Conditional-formatting fills need both start_color and end_color set, or Excel's
# differential styling won't actually render the background -- fgColor alone is silently
# ignored in this context, unlike a normal (non-conditional) cell fill.
RED_HIGHLIGHT = {
	"fill": PatternFill(fill_type="solid", start_color="FF0000", end_color="FF0000"),
	"font": Font(color="FFFFFF"),
}
ORANGE_HIGHLIGHT = {
	"fill": PatternFill(fill_type="solid", start_color="FFA500", end_color="FFA500"),
	"font": Font(color="000000"),
}


def _add_summary_block(sheet, header, last_data_row):
	# A small label/value table below the data, using live SUM formulas (not pre-computed
	# numbers) so it stays correct if someone edits the sheet by hand. Excel's SUM() skips
	# non-numeric cells on its own, so "CROSSLISTED" entries in these columns are ignored.
	# Only added once ta_assignment/ta_needs_estimation.py has filled these columns in.
	if RECITATION_COLUMN not in header or GRADER_COLUMN not in header:
		return

	recitation_letter = get_column_letter(header.index(RECITATION_COLUMN) + 1)
	grader_letter = get_column_letter(header.index(GRADER_COLUMN) + 1)

	recitation_row = last_data_row + 2  # one blank row gap below the data
	grading_row = recitation_row + 1
	total_row = grading_row + 1

	sheet.cell(row=recitation_row, column=1, value="Total Recitation")
	sheet.cell(row=recitation_row, column=2, value=f"=SUM({recitation_letter}2:{recitation_letter}{last_data_row})")

	sheet.cell(row=grading_row, column=1, value="Total Grading")
	sheet.cell(row=grading_row, column=2, value=f"=SUM({grader_letter}2:{grader_letter}{last_data_row})")

	sheet.cell(row=total_row, column=1, value="Total TAs")
	sheet.cell(row=total_row, column=2, value=f"=B{recitation_row}+B{grading_row}")


def _add_enrollment_highlighting(sheet, header, last_data_row):
	# Flags low-enrollment lecture/PRA/INT rows so a reviewer can spot them at a glance:
	# red under 20 students, orange from 20-24. Live conditional-formatting rules (not
	# colors baked in at generation time), so they stay correct if enrollment is edited by
	# hand later. The two ranges are written as mutually exclusive so a <20 row always
	# reads red, never orange.
	if TYPE_COLUMN not in header or ENROLLMENT_COLUMN not in header:
		return

	type_letter = get_column_letter(header.index(TYPE_COLUMN) + 1)
	enrollment_letter = get_column_letter(header.index(ENROLLMENT_COLUMN) + 1)
	full_row_range = f"A2:{get_column_letter(len(header))}{last_data_row}"

	is_primary_lecture = (
		f'OR(${type_letter}2="LEC",${type_letter}2="PRA",${type_letter}2="INT")'
	)
	is_real_number = f"ISNUMBER(${enrollment_letter}2)"

	red_condition = f"AND({is_primary_lecture},{is_real_number},${enrollment_letter}2<{RED_ENROLLMENT_CUTOFF})"
	orange_condition = (
		f"AND({is_primary_lecture},{is_real_number},"
		f"${enrollment_letter}2>={RED_ENROLLMENT_CUTOFF},${enrollment_letter}2<{ORANGE_ENROLLMENT_CUTOFF})"
	)

	sheet.conditional_formatting.add(full_row_range, FormulaRule(formula=[red_condition], **RED_HIGHLIGHT))
	sheet.conditional_formatting.add(full_row_range, FormulaRule(formula=[orange_condition], **ORANGE_HIGHLIGHT))


def _as_number(value):
	# Leaves non-numeric text (e.g. "CROSSLISTED") and blanks untouched, so they still
	# render as-is; only genuinely numeric strings become real numbers.
	if value == "":
		return value
	try:
		return int(value)
	except ValueError:
		pass
	try:
		return float(value)
	except ValueError:
		return value


def gen_xlsx(csv_path: str, output_path: str = None) -> str:
	# Converts a class report CSV to XLSX. If the CSV already has TA-needs columns, adds a
	# small summary block below the data with live totals.
	if output_path is None:
		base_name = os.path.splitext(os.path.basename(csv_path))[0]
		output_path = os.path.join(REPORTS_DIR, f"{base_name}.xlsx")

	with open(csv_path, newline="") as f:
		header, *data_rows = csv.reader(f)

	numeric_cols = [header.index(column) for column in NUMERIC_COLUMNS if column in header]

	workbook = Workbook()
	sheet = workbook.active
	sheet.append(header)
	for row in data_rows:
		row = list(row)
		for col in numeric_cols:
			row[col] = _as_number(row[col])
		sheet.append(row)

	last_data_row = len(data_rows) + 1  # +1 to account for the header row
	
	_add_summary_block(sheet, header, last_data_row)
	_add_enrollment_highlighting(sheet, header, last_data_row)

	workbook.save(output_path)
	return output_path
